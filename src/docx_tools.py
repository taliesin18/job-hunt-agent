"""
Utilities for editing .docx files in place without breaking formatting.

Works with a "placeholder marker" template: you tag your real resume
once with tokens like {{SUMMARY}} or {{EXP1_BULLET2}}, and these
functions replace only the text inside those tokens — never deleting
or creating runs — so fonts, bold job titles, bullet styles, tables,
and spacing are all left exactly as they were.

Why this matters: Word stores paragraph text as a sequence of "runs,"
each carrying its own formatting. Setting `paragraph.text = ...`
destroys every run and replaces the whole paragraph with one plain
run, which is what breaks formatting. This module only ever mutates
`run.text` on the specific run(s) a token spans, so every other
character's formatting object is untouched.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

_TOKEN_PATTERN = re.compile(r"\{\{[A-Z0-9_]+\}\}")


def _iter_table_paragraphs(table: Table) -> list[Paragraph]:
    paragraphs: list[Paragraph] = []
    for row in table.rows:
        for cell in row.cells:
            paragraphs.extend(cell.paragraphs)
            for nested in cell.tables:
                paragraphs.extend(_iter_table_paragraphs(nested))
    return paragraphs


def _iter_all_paragraphs(document) -> list[Paragraph]:
    """Body paragraphs plus every paragraph inside every table cell
    (including nested tables) — resumes often use tables for layout."""
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        paragraphs.extend(_iter_table_paragraphs(table))
    return paragraphs


def find_placeholders(docx_path: str) -> list[str]:
    """Return every {{TOKEN}} found anywhere in the document, deduplicated
    and sorted."""
    document = Document(docx_path)
    tokens: set[str] = set()
    for paragraph in _iter_all_paragraphs(document):
        full_text = "".join(run.text for run in paragraph.runs)
        tokens.update(_TOKEN_PATTERN.findall(full_text))
    return sorted(tokens)


def _replace_first_occurrence(paragraph: Paragraph, token: str, replacement: str) -> bool:
    """Replace one occurrence of `token` inside `paragraph`, editing only
    the run(s) it spans. Returns True if a replacement was made."""
    runs = paragraph.runs
    full_text = "".join(r.text for r in runs)
    idx = full_text.find(token)
    if idx == -1:
        return False
    end = idx + len(token)

    offsets = []
    pos = 0
    for r in runs:
        offsets.append((pos, pos + len(r.text)))
        pos += len(r.text)

    touched = [i for i, (s, e) in enumerate(offsets) if e > idx and s < end]
    if not touched:
        return False

    first_i, last_i = touched[0], touched[-1]
    prefix = full_text[offsets[first_i][0]:idx]
    suffix = full_text[end:offsets[last_i][1]]

    if first_i == last_i:
        runs[first_i].text = prefix + replacement + suffix
    else:
        # Replacement text and the prefix inherit the first run's
        # formatting; any text after the token inherits the last run's.
        runs[first_i].text = prefix + replacement
        runs[last_i].text = suffix
        for i in touched[1:-1]:
            runs[i].text = ""
    return True


def fill_template(docx_path: str, replacements: dict[str, str], output_path: str) -> dict:
    """
    Replace every {{TOKEN}} in the document with replacements[TOKEN]
    (keys may be given with or without the surrounding braces), and
    save the result to output_path. The original file is never
    modified in place.

    Returns a report so a mismatch is visible instead of silently
    producing a half-filled document:
        {
          "filled": [tokens successfully replaced],
          "not_found_in_doc": [replacement keys that had no matching token],
          "not_provided": [tokens in the doc that got no replacement],
        }
    """
    document = Document(docx_path)
    doc_tokens = set(find_placeholders(docx_path))

    filled: list[str] = []
    for key, value in replacements.items():
        plain = key.strip("{}")
        full_token = f"{{{{{plain}}}}}"
        found_anywhere = False
        for paragraph in _iter_all_paragraphs(document):
            while _replace_first_occurrence(paragraph, full_token, value):
                found_anywhere = True
        if found_anywhere:
            filled.append(plain)

    normalized_doc_tokens = {t.strip("{}") for t in doc_tokens}
    normalized_provided = {k.strip("{}") for k in replacements.keys()}

    not_found_in_doc = sorted(normalized_provided - normalized_doc_tokens)
    not_provided = sorted(normalized_doc_tokens - normalized_provided)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)

    return {
        "filled": sorted(filled),
        "not_found_in_doc": not_found_in_doc,
        "not_provided": not_provided,
    }
