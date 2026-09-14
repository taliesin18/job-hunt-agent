"""
Utilities for the Excel-based resume workflow.

Design: the workbook has two sheets —
  - "Data": one row per field (column A = field name, column B = value).
    This is the only sheet these functions ever write to.
  - "Resume": a printable layout whose cells are formulas like
    =Data!B5, referencing the Data sheet. Never touched here — Excel's
    own formula engine is what "preserves formatting," so there's no
    run-splitting logic to get right (unlike docx_tools.py).

This sidesteps the docx approach's main source of fragility: instead
of finding and replacing text inside Word's run structure, we only
ever write plain values into designated cells, and the spreadsheet's
own formulas do the formatting-safe part for free.
"""
from __future__ import annotations

import csv
import shutil
import subprocess
from pathlib import Path

import openpyxl

DATA_SHEET_NAME = "Data"
FIELD_COL = 1   # column A
VALUE_COL = 2   # column B
HEADER_ROW = 1  # row 1 is the "Field" / "Value" header


def read_data_fields(xlsx_path: str) -> list[str]:
    """Return every field name in the Data sheet's column A (skipping the
    header row), in row order."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    if DATA_SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"No '{DATA_SHEET_NAME}' sheet found in {xlsx_path}.")
    sheet = wb[DATA_SHEET_NAME]
    fields = []
    for row in sheet.iter_rows(min_row=HEADER_ROW + 1, min_col=FIELD_COL, max_col=FIELD_COL):
        value = row[0].value
        if value:
            fields.append(str(value).strip())
    return fields


def read_data_field_values(xlsx_path: str) -> dict:
    """Return {field_name: current_value} from the Data sheet — used to
    ground the LLM on what each field is currently about (see
    field_filler.generate_field_values' field_context parameter), since
    unlike the docx template this sheet's original content is never
    blanked out."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    if DATA_SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"No '{DATA_SHEET_NAME}' sheet found in {xlsx_path}.")
    sheet = wb[DATA_SHEET_NAME]
    values = {}
    for row in sheet.iter_rows(min_row=HEADER_ROW + 1, min_col=FIELD_COL, max_col=VALUE_COL):
        field_cell, value_cell = row[0], row[1]
        if field_cell.value:
            values[str(field_cell.value).strip()] = str(value_cell.value or "")
    return values


def write_data_values(xlsx_path: str, values: dict, output_path: str) -> dict:
    """
    Write values[field] into column B of the Data sheet row where column A
    == field, for every field found in both the sheet and `values`. Saves
    to output_path (never modifies xlsx_path in place). The Resume sheet's
    formulas are untouched — they'll show the new values next time the
    workbook is opened/recalculated.

    Returns {"filled": [...], "not_found_in_sheet": [...], "not_provided": [...]}
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    if DATA_SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"No '{DATA_SHEET_NAME}' sheet found in {xlsx_path}.")
    sheet = wb[DATA_SHEET_NAME]

    sheet_fields: dict[str, int] = {}
    for row in sheet.iter_rows(min_row=HEADER_ROW + 1, min_col=FIELD_COL, max_col=FIELD_COL):
        cell = row[0]
        if cell.value:
            sheet_fields[str(cell.value).strip()] = cell.row

    filled = []
    for field, value in values.items():
        row_num = sheet_fields.get(field)
        if row_num is not None:
            sheet.cell(row=row_num, column=VALUE_COL, value=value)
            filled.append(field)

    not_found_in_sheet = sorted(set(values.keys()) - set(filled))
    not_provided = sorted(set(sheet_fields.keys()) - set(values.keys()))

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    return {
        "filled": sorted(filled),
        "not_found_in_sheet": not_found_in_sheet,
        "not_provided": not_provided,
    }


def export_to_csv(values: dict, csv_path: str) -> None:
    """Write {field: value} as a two-column field,value CSV — useful as a
    plain, portable record of what the resume-fixer produced, independent
    of Excel."""
    Path(csv_path).parent.mkdir(parents=True, exist_ok=True)
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Field", "Value"])
        for field, value in values.items():
            writer.writerow([field, value])


def convert_to_pdf(xlsx_path: str, output_dir: str | None = None) -> str:
    """
    Export xlsx_path to PDF via LibreOffice headless, respecting each
    sheet's defined print area (hidden sheets, like Data, are skipped
    automatically). Returns the PDF's path.
    """
    if shutil.which("soffice") is None:
        raise RuntimeError(
            "LibreOffice ('soffice') was not found on PATH. Install LibreOffice "
            "to use PDF export, or open the .xlsx in Excel and print to PDF manually."
        )
    out_dir = output_dir or str(Path(xlsx_path).parent)
    Path(out_dir).mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", out_dir, xlsx_path],
        check=True, capture_output=True,
    )
    pdf_path = Path(out_dir) / (Path(xlsx_path).stem + ".pdf")
    if not pdf_path.exists():
        raise RuntimeError(f"Expected PDF at {pdf_path} but it wasn't created.")
    return str(pdf_path)
